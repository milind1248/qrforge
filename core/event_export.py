"""Excel reports for events: full attendee list, actual check-in records, scan log and summary (openpyxl)."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core import event_db as E

HEAD_FILL = PatternFill("solid", fgColor="4F46E5")
STATUS_LABEL = {"confirmed": "Confirmed", "pending_payment": "Payment pending", "cancelled": "Cancelled", "rejected": "Payment rejected"}


def _put(ws, row: list, header: bool = False):
    r = getattr(ws, "_next_row", 0) + 1
    ws._next_row = r
    for i, v in enumerate(row, 1):
        c = ws.cell(row=r, column=i)
        if isinstance(v, str):
            c.value = v
            c.data_type = "s"            # always text: a name like =HYPERLINK(...) must never run as a formula
        else:
            c.value = v
        if header:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = HEAD_FILL
            c.alignment = Alignment(vertical="center")


def _fit(ws):
    for col in ws.columns:
        w = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(10, w + 2), 48)
    ws.freeze_panes = "A2"


def _ist(utc):
    d = E.utc_to_ist(utc) if utc else None
    return d.strftime("%Y-%m-%d %H:%M:%S") if d else ""


def _attendee_row(b: dict) -> list:
    return [b["event_name"], b["ref"], b["name"], b["email"], b["phone"] or "", b["extra"] or "", STATUS_LABEL.get(b["status"], b["status"]),
            b["amount"] or 0, b["pay_ref"] or "", "Yes" if b["checked_in_at"] and b["status"] == "confirmed" else "No",
            _ist(b["checked_in_at"]), b["checked_in_by"] or "", _ist(b["created_at"]), _ist(b["cancelled_at"]), b["review_note"] or ""]


ATT_HEAD = ["Event", "Booking ID", "Name", "Email", "Phone", "Extra detail", "Status", "Amount (INR)", "UPI reference", "Checked in",
            "Check-in time (IST)", "Checked in by", "Booked at (IST)", "Cancelled at (IST)", "Review note"]


def build(events: list[dict], rows: list[dict], kind: str = "full", include_scan_log: bool = True) -> bytes:
    """kind='full'     -> every booking in `rows` (the complete attendee list)
       kind='checkins' -> only attendees who actually checked in (one row per check-in, with time)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    _put(ws, ["Event", "Date", "Time", "Venue", "Capacity", "Confirmed", "Payment pending", "Cancelled/rejected", "Checked in", "Attendance %", "Revenue (INR)"], True)
    for ev in events:
        s = E.stats(ev["id"])
        _put(ws, [ev["name"], ev["event_date"], f"{ev['start_time']}" + (f" - {ev['end_time']}" if ev.get("end_time") else ""), ev["venue"], ev["capacity"],
                  s["confirmed"], s["pending"], s["cancelled"], s["checked_in"], s["rate"], s["revenue"]])
    _fit(ws)

    if kind == "checkins":
        chk = sorted([b for b in rows if b["checked_in_at"] and b["status"] == "confirmed"], key=lambda b: b["checked_in_at"])
        ws2 = wb.create_sheet("Check-ins")
        _put(ws2, ["#", "Event", "Booking ID", "Name", "Email", "Phone", "Check-in date (IST)", "Check-in time (IST)", "Checked in by"], True)
        for i, b in enumerate(chk, 1):
            t = _ist(b["checked_in_at"])
            _put(ws2, [i, b["event_name"], b["ref"], b["name"], b["email"], b["phone"] or "", t[:10], t[11:], b["checked_in_by"] or ""])
        _fit(ws2)
    else:
        ws2 = wb.create_sheet("Attendees")
        _put(ws2, ATT_HEAD, True)
        for b in rows:
            _put(ws2, _attendee_row(b))
        _fit(ws2)
        ws2.auto_filter.ref = ws2.dimensions

    if include_scan_log:
        ws3 = wb.create_sheet("Scan log")
        _put(ws3, ["Event", "Time (IST)", "Result", "Booking ID", "Name", "Scanned by"], True)
        for ev in events:
            for r in E._q("SELECT s.ts, s.result, s.actor, b.ref, b.name FROM pass_scans s LEFT JOIN bookings b ON b.id=s.booking_id"
                          " WHERE s.event_id=? AND s.result<>'bad_pin' ORDER BY s.id", (ev["id"],)):
                _put(ws3, [ev["name"], _ist(r["ts"]), r["result"], r["ref"] or "", r["name"] or "", r["actor"] or ""])
        _fit(ws3)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
