import io
from datetime import timedelta

import openpyxl
import pytest

from core import auth, event_db as E, event_export, event_mail, event_pass as P


def _data(**kw):
    d = (E.now_ist() + timedelta(days=5)).strftime("%Y-%m-%d")
    base = dict(name="Expo", description="", venue="Hall", map_url="", event_date=d, start_time="10:00", end_time="12:00", capacity=20, reg_deadline="",
                max_tickets=5, is_paid=0, price=0, upi_id="", upi_name="", pay_note="", collect_phone=1, custom_label="", contact_info="")
    base.update(kw)
    return base


@pytest.fixture()
def setup(fresh_db):
    uid = fresh_db.create_user("o@test.local", "Org", auth.hash_pw("x" * 10))
    ev = E.create_event(uid, _data())
    E.create_bookings(ev, ["Asha", "=HYPERLINK(\"http://evil\",\"x\")"], "a@x.com", "9999999999")
    E.create_bookings(ev, ["Ben"], "b@x.com")
    E.create_bookings(ev, ["Cara"], "c@x.com")
    return ev


def _wb(data):
    return openpyxl.load_workbook(io.BytesIO(data))


def test_full_export_has_headers_rows_and_safe_text(setup):
    ev = setup
    rows = E.search_bookings([ev["id"]])
    E.process_scan(ev, rows[-1]["ref"], "gate")                     # Asha checks in
    E.cancel([rows[0]["id"]], "org")                                 # Cara cancelled
    wb = _wb(event_export.build([ev], E.search_bookings([ev["id"]]), "full"))
    assert wb.sheetnames == ["Summary", "Attendees", "Scan log"]
    s = [[c.value for c in r] for r in wb["Summary"]]
    assert s[0][0] == "Event" and s[1][0] == "Expo" and s[1][4] == 20 and s[1][5] == 3 and s[1][8] == 1       # header + capacity + confirmed + checked in
    a = [[c.value for c in r] for r in wb["Attendees"]]
    assert a[0][:3] == ["Event", "Booking ID", "Name"] and len(a) == 5                                         # header + 4 bookings
    names = {r[2] for r in a[1:]}
    assert '=HYPERLINK("http://evil","x")' in names
    cell = next(c for r in wb["Attendees"] for c in r if c.value and str(c.value).startswith("=HYPERLINK"))
    assert cell.data_type == "s"                                                                              # stored as text, never a formula
    by = {r[2]: r for r in a[1:]}
    assert by["Asha"][9] == "Yes" and by["Asha"][10] and by["Ben"][9] == "No" and by["Cara"][6] == "Cancelled"
    log = [[c.value for c in r] for r in wb["Scan log"]]
    assert log[0][0] == "Event" and log[1][2] == "valid" and log[1][5] == "gate"


def test_checkin_only_export_and_time_filter(setup):
    ev = setup
    rows = E.search_bookings([ev["id"]])
    E.process_scan(ev, rows[-1]["ref"], "gate")
    E.process_scan(ev, rows[-2]["ref"], "gate")
    ci = [b for b in E.search_bookings([ev["id"]], checked="in") if b["status"] == "confirmed"]
    wb = _wb(event_export.build([ev], ci, "checkins", include_scan_log=False))
    assert wb.sheetnames == ["Summary", "Check-ins"]
    r = [[c.value for c in x] for x in wb["Check-ins"]]
    assert r[0][:4] == ["#", "Event", "Booking ID", "Name"] and len(r) == 3 and r[1][6] and r[1][7] and r[1][8] == "gate"
    # a time window in the past returns nobody, a window around now returns both
    assert E.search_bookings([ev["id"]], checked="in", ci_to=(E.now_ist() - timedelta(days=2)).strftime("%Y-%m-%d %H:%M:%S")) == []
    assert len(E.search_bookings([ev["id"]], checked="in", ci_from="2000-01-01 00:00:00", ci_to="2100-01-01 00:00:00")) == 2


def test_email_builders_do_nothing_when_email_disabled(setup):
    ev = setup
    bks = E.search_bookings([ev["id"]])
    full = [E.get_booking(b["id"]) for b in bks]
    # QRFORGE_NO_EMAIL=1 -> notify.enabled() is False -> none of these may raise or try SMTP
    event_mail.send_passes(ev, full)
    event_mail.send_payment_received(ev, full)
    event_mail.send_rejected(ev, full[0], "x")
    event_mail.send_cancelled(ev, full[0])
    event_mail.notify_organizer_payment(ev, {"email": "o@test.local"}, full)


def test_pass_artifacts_decode(setup):
    from core import qr_engine
    ev = setup
    b = E.get_booking(E.search_bookings([ev["id"]])[0]["id"])
    png = P.pass_png(ev, b)
    assert png[:8] == b"\x89PNG\r\n\x1a\n"
    assert qr_engine.decode(png) == P.pass_url(b["token"])
    assert P.pass_pdf(ev, b)[:4] == b"%PDF"
    ics = P.ics_text(ev, b)
    assert ics.startswith("BEGIN:VCALENDAR") and b["ref"] in ics and "DTSTART:" in ics
    assert "wa.me/919999999999" in P.whatsapp_link("9999999999", "hi")
